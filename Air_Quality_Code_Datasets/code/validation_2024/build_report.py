from pathlib import Path
import json,pandas as pd,zipfile,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parent;base=root.parent
scores=pd.read_csv(root/'comparison_2024.csv').set_index('method');old=pd.read_csv(base/'source/pkg/timexer_results/patch_6/comparison_all_methods.csv').set_index('method');d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(22);d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 2'].font.size=Pt(11)
footer=sec.footer.paragraphs[0];footer.text='PM2.5 frozen model cross year evaluation  |  Abhishek Patil  |  ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
labels={'Hour median':'Hour of day median','Linear interpolation':'Linear interpolation','Random Forest':'Random Forest','PM2.5 station-temporal Transformer':'Station temporal Transformer','TimeXer imputation adaptation':'TimeXer imputation adaptation'}
order=['Hour median','Linear interpolation','Random Forest','PM2.5 station-temporal Transformer','TimeXer imputation adaptation']
def p(t):d.add_paragraph(t)
def h(t):d.add_heading(t,1)
def table(headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
 for c,v in zip(t.rows[0].cells,headers):
  c.text=str(v);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');c._tc.get_or_add_tcPr().append(sh)
  for run in c.paragraphs[0].runs:run.bold=True
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(3);par.paragraph_format.space_before=Pt(3)
    for r in par.runs:r.font.size=Pt(9)
 if widths:
  t.autofit=False
  for row in t.rows:
   for c,w in zip(row.cells,widths):c.width=Inches(w)
 return t

d.add_paragraph('PM25 Cross Year Validation Results',style='Title');p('Prepared for Professor Ke Yang by Abhishek Patil | 30 September 2026')
p('I evaluated the frozen 2022 models on 2024 observations from the same three PM2.5 stations. Random Forest performed best on this cross-year test. The station-temporal Transformer still had a lower pooled error than interpolation, but its October validation advantage over Random Forest did not generalize to the new year. The TimeXer imputation adaptation had higher error than both methods.')
h('Results in original concentration units')
table(['Method','October 2022\nMAE / RMSE','2024\nMAE / RMSE'],[[labels[k],f'{old.loc[k,"MAE"]:.3f} / {old.loc[k,"RMSE"]:.3f}',f'{scores.loc[k,"MAE"]:.3f} / {scores.loc[k,"RMSE"]:.3f}'] for k in order],[3.15,1.6,1.6])
p('All errors are in µg/m³. October contains 886 hidden seed–timestamp–station instances; 2024 contains 10,305. Every method uses the same targets within each evaluation. These counts include repeated observations across masking seeds. October was used for model selection; 2024 was evaluated after the current checkpoints and settings were frozen. This is a same-region transfer test, not a test of generalization to a new region.')
h('What the uncertainty shows')
p('The 2024 Transformer MAE is 0.062 µg/m³ higher than Random Forest and 0.131 lower than interpolation. A paired seven-day block bootstrap gives a 95% interval of approximately 0.000 to 0.122 for Transformer minus Random Forest and −0.248 to 0.004 for Transformer minus interpolation. The first interval is close to zero; the second includes zero. These estimates describe uncertainty in this year of evaluation, not variation across fresh training runs.')
h('Checks completed')
p('I loaded the saved checkpoints and independently reproduced all five October methods. The largest neural prediction discrepancy was 0.0000034 µg/m³, consistent with floating-point inference variation. Baseline discrepancies were below 0.00000001. I verified training-only scalers, identical evaluation keys, removal of hidden targets from neural inputs, finite predictions and chronological separation. No model was retrained or tuned on 2024.')

h('Data sources and quality screening');d.paragraphs[-1].paragraph_format.page_break_before=True
p('The 2024 targets come from the official EPA AQS hourly archive for parameter 88101, using Texas state 48, Bexar County 029, sites 0032, 0059 and 1069, and POC 2. The method code is 209 at all three sites, matching the 2022 hourly records. Daily POC 1 measurements are excluded. EPA GMT timestamps are used directly; every selected record has a six-hour local-standard-to-GMT offset, and there are no duplicate station-hours.')
p('I retained finite numbers with an empty EPA qualifier, matching the pilot’s observed_valid rule. That label is a screening rule, not independent certification of every measurement. There are 24,794 screened observations out of 26,352 possible station-hours. The UTC calendar year has 8,784 hours, including leap day; the 2022 pilot instead used a CST calendar-year index. Evaluation windows start at 00:00 UTC in the 2024 test.')
p('The twelve original monthly ERA5 files supply seven complete hourly fields at latitude 29.5° N, longitude 98.5° W: u10, v10, t2m, d2m, sp, tp and blh. These are shared single-grid-point inputs, not distinct weather at each station. Source units are m/s, K, Pa, m and m for precipitation depth. Source UTC timestamps and the coordinates were checked directly. The current 2022 pilot weather matches the previously source-verified series within rounding tolerances at all 8,760 hours. The earlier November daylight-saving correction is already reflected in this pilot; it does not change training or October results.')
h('Negative readings require a separate interpretation')
p('The 2024 screened data contain 983 negative readings; 66 lie below the reported negative MDL threshold of −5 µg/m³. Negative numbers were retained in the primary comparison to preserve the frozen pilot rule. The 2022 source audit also contains below-negative-MDL observations. These records require measurement-quality review before a publication dataset is frozen. Excluding the 10 scored 2024 instances below −5 gives MAEs of 2.752 for Random Forest, 2.815 for the station Transformer and 3.960 for TimeXer; the ranking is unchanged. This sensitivity removes scored targets only; it does not clean model inputs or retrain models.')
h('Training and features remained fixed')
table(['Method','Features and fitting'],[
['Hour median','Station and local-standard hour; medians fitted on January–September 2022.'],
['Interpolation','Visible values immediately surrounding each artificially hidden gap; no fitting.'],
['Random Forest','Visible PM2.5 with ±1, ±2, ±3, ±6 and ±24 hour context and other stations at the target hour; weather, hour and day-of-year encodings, station identity. 160 trees; depth 16; leaf size 3; training-fitted feature median.'],
['Station Transformer','24 hourly tokens, each with 3 masked PM2.5 values, 3 visibility flags, 7 weather fields and 2 hour encodings. 64 dimensions, 4 heads, 2 layers; best epoch 16.'],
['TimeXer adaptation','Official TimeXer patch embedding and encoder; 3 endogenous station series, 15 masked exogenous channels; same-window reconstruction head. Frozen patch length 6; best epoch 9.']],[1.35,5.0])
p('Both neural models were trained on January–September 2022 with train-only standardization, new observed-value masks each epoch, masked MAE loss, AdamW learning rate 0.001, batch size 32 and early stopping selected by the three October masks. Training used one initialization seed, 172. TimeXer patch lengths 3 and 4 were previously explored on October; the patch-6 checkpoint was frozen before this test.')

h('Performance by station and gap length');d.paragraphs[-1].paragraph_format.page_break_before=True
station=pd.read_csv(root/'by_station.csv').set_index(['method','site']);gap=pd.read_csv(root/'by_gap.csv').set_index(['method','gap_hours'])
table(['Method','Site 0032','Site 0059','Site 1069'],[[labels[k]]+[f'{station.loc[(k,s),"MAE"]:.3f}' for s in [32,59,1069]] for k in order],[3.05,1.1,1.1,1.1])
p('Station entries are MAE in µg/m³. Random Forest performs best at 0032 and 0059; the station Transformer performs best at 1069. A pooled ranking therefore does not describe every monitor.')
table(['Method','1 hour gaps','3 hour gaps','6 hour gaps'],[[labels[k]]+[f'{gap.loc[(k,g),"MAE"]:.3f}' for g in [1,3,6]] for k in order],[3.05,1.1,1.1,1.1])
p('Gap entries are MAE. There are 1,113 scored instances for one-hour gaps, 3,144 for three-hour gaps and 6,048 for six-hour gaps. Interpolation performs best for one-hour gaps; Random Forest performs best for three- and six-hour gaps. Longer gaps contribute more scored hourly targets, so pooled MAE is not an equal-weight average of gap types.')
h('Three concrete examples')
e=pd.read_csv(root/'examples_2024.csv')
table(['UTC time at site 0059','Gap','Truth','RF','Station\nTransformer','TimeXer'],[[x.datetime_utc[:16].replace('T',' '),str(x.gap_hours)+' h',f'{x.truth:.2f}',f'{x["Random Forest"]:.2f}',f'{x["PM2.5 station-temporal Transformer"]:.2f}',f'{x["TimeXer imputation adaptation"]:.2f}'] for _,x in e.iterrows()],[1.7,.55,.6,.7,1.35,1.05])
p('These are the earliest hidden timestamps at site 0059 for each gap length under seed 11, chosen without searching for favorable errors. For the one-hour example, the Transformer estimates 8.61 against a truth of 8.40. For the three-hour example, Random Forest is closer than either neural model. For the six-hour example, the Transformer is closer than Random Forest and TimeXer. A single example illustrates behavior; aggregate scores establish the comparison.')
h('What is now validated and what comes next');d.paragraphs[-1].paragraph_format.page_break_before=True
p('The saved models, reported October predictions, train-only scalers, source weather alignment and frozen-model 2024 short-gap comparison have been checked. The source files, masks, individual predictions, hashes, station/gap/seed summaries and bootstrap outputs are included in the companion package. Its runners reproduce data preparation, checkpoint inference and metric calculation.')
p('The next research step is to review below-detection-limit records and freeze a final quality policy; then repeat training across several initialization seeds and test natural-outage-like masks and another region. Artificial gaps require valid endpoints and cover only 1, 3 and 6 hours. The year 2024 appeared in earlier exploratory project work, so it should not be described as never previously inspected. Its scores were not used to fit or select these frozen models.')
p('A separate forecasting experiment is still needed to answer Professor Ke’s original TimeXer forecasting question. This comparison uses an imputation adaptation of the official code, not the unchanged forecasting model. Both interpolation and the neural imputers use later observations within the reconstruction window, and ERA5 is retrospective reanalysis. A forecasting benchmark must enforce information available at each forecast origin and use weather inputs available then. No forecasting or publication-wide validation claim follows from this experiment.')
p('Sources: EPA hourly archive https://aqs.epa.gov/aqsweb/airdata/hourly_88101_2024.zip; official TimeXer https://github.com/thuml/TimeXer at commit 76011909357972bd55a27adba2e1be994d81b327. ERA5 source files were supplied with the project.')

for node in d.element.xpath('//w:pBdr'):
 node.getparent().remove(node)
for style in d.styles:
 for node in style.element.xpath('.//w:pBdr'):node.getparent().remove(node)
report=base/'PM25_Cross_Year_Validation_Report.docx';d.save(report)
(root/'README.md').write_text('''# PM2.5 frozen-model 2024 transfer evaluation

Primary result: Random Forest MAE 2.763, station-temporal Transformer 2.826, interpolation 2.957, TimeXer adaptation 3.969 µg/m³. Do not carry forward the October claim that the Transformer beats Random Forest to the 2024 evaluation.

Run prepare_data.py from this directory after downloading the official archive to raw/hourly_88101_2024.zip and providing the twelve original ERA5 monthly files under ../validation_sources/NEW_NASA_RESEARCH. The original 2022 ERA5 canonical CSV is provided in validation_sources. prepare_data.py checks source alignment and creates panel_2024.csv. The included panel enables checkpoint evaluation without a new source download.

Install CPU torch 2.6.0, numpy, pandas, scikit-learn 1.8.0, joblib, einops and reformer-pytorch and its import dependencies. Run python evaluate.py followed by python additional_checks.py. Frozen 2022 code/checkpoints/data are in source/pkg. evaluate.py first replays October neural predictions, verifies training-only scalers, then evaluates 2024. additional_checks.py replays all three baselines and computes weekly-block uncertainty. Primary daily bootstrap seed404; weekly sensitivity seed407; 2000 replicates. Dependence across masking seeds is grouped by UTC day. No retraining or parameter selection is performed.

The annual EPA national archive is omitted to keep the package compact. Its URL and SHA256 are in data_audit.json; raw/aqs_2024_bexar_original.csv contains the original filtered records. Weather values are decoded from original source monthly files and stored with attributes and timestamps. Negative values follow the original screening rule; target_quality_sensitivity.csv reports exclusions from scoring only. Training data and visible inputs are not cleaned by that sensitivity. No forecasts were run.
''')
print(report)
